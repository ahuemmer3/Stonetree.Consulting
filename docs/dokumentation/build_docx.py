"""Erzeugt die Projektdokumentation als Word-Datei (und, falls Word vorhanden ist, als PDF).

Quelle ist projektdokumentation.md im selben Ordner. Unterstützt wird nur das,
was die Dokumentation braucht:

    # Kapitel / ## Abschnitt     nummeriert, mit {-} am Ende unnummeriert
    ![Beschriftung](pfad)         Abbildung mit Beschriftung darunter
    Tabelle: Beschriftung         direkt vor einer Markdown-Tabelle
    - Punkt / 1. Punkt            Aufzählungen
    **fett**, `code`              Hervorhebungen im Text
    <!-- toc|lof|lot -->         Inhalts-, Abbildungs-, Tabellenverzeichnis
    <!-- frontmatter -->         ab hier römische Seitenzahlen
    <!-- main -->                ab hier arabische Seitenzahlen ab 1
    <!-- quellen -->             folgende [n]-Absätze mit hängendem Einzug

Aufruf:  python docs/dokumentation/build_docx.py [zieldatei.docx]
Benötigt: pip install python-docx pillow
"""

import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HIER = Path(__file__).resolve().parent
QUELLE = HIER / "projektdokumentation.md"
ZIEL = HIER.parent / "Projektdokumentation-stonetree.docx"

SCHRIFT = "Arial"
TEXT = RGBColor(0x12, 0x1A, 0x1F)
GEDAEMPFT = RGBColor(0x5F, 0x6C, 0x75)
KOPF_HINTERGRUND = "E3E9EC"
LINIE = "C9CFD4"
TEXTBREITE_CM = 15.5
MAX_BILDHOEHE_CM = 12.0


# ---------------------------------------------------------------- Hilfen für XML


def feld(absatz, anweisung: str, platzhalter: str = "") -> None:
    """Fügt ein Word-Feld ein (Seitenzahl, Verzeichnis, Nummerierung)."""

    def zeichen(typ):
        lauf = absatz.add_run()
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), typ)
        lauf._r.append(element)

    zeichen("begin")
    lauf = absatz.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {anweisung} "
    lauf._r.append(instr)
    zeichen("separate")
    absatz.add_run(platzhalter)
    zeichen("end")


def zellfarbe(zelle, farbe: str) -> None:
    schattierung = OxmlElement("w:shd")
    schattierung.set(qn("w:val"), "clear")
    schattierung.set(qn("w:color"), "auto")
    schattierung.set(qn("w:fill"), farbe)
    zelle._tc.get_or_add_tcPr().append(schattierung)


def tabellenrahmen(tabelle) -> None:
    rahmen = OxmlElement("w:tblBorders")
    for seite in ("top", "left", "bottom", "right", "insideH", "insideV"):
        linie = OxmlElement(f"w:{seite}")
        linie.set(qn("w:val"), "single")
        linie.set(qn("w:sz"), "4")
        linie.set(qn("w:color"), LINIE)
        rahmen.append(linie)
    eigenschaften = tabelle._tbl.tblPr
    eigenschaften.append(rahmen)
    breite = OxmlElement("w:tblW")
    breite.set(qn("w:type"), "pct")
    breite.set(qn("w:w"), "5000")
    for alt in eigenschaften.findall(qn("w:tblW")):
        eigenschaften.remove(alt)
    eigenschaften.append(breite)


def zeile_zusammenhalten(zeile) -> None:
    element = OxmlElement("w:cantSplit")
    element.set(qn("w:val"), "true")
    zeile._tr.get_or_add_trPr().append(element)


def kopfzeile_wiederholen(zeile) -> None:
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    zeile._tr.get_or_add_trPr().append(element)


def seitenzahlen(abschnitt, format_: str, start: int = 1) -> None:
    element = OxmlElement("w:pgNumType")
    element.set(qn("w:fmt"), format_)
    element.set(qn("w:start"), str(start))
    abschnitt._sectPr.append(element)


def nicht_trennen(absatz) -> None:
    absatz.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------- Formatvorlagen


def vorlagen(doc) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = SCHRIFT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SCHRIFT)
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(6)

    for name, groesse, vorher, nachher in (
        ("Heading 1", 16, 0, 12),
        ("Heading 2", 13, 14, 6),
        ("Title", 24, 0, 6),
    ):
        stil = doc.styles[name]
        stil.font.name = SCHRIFT
        # Designschrift entfernen, sonst überschreibt sie die Schriftart
        schriften = stil.element.rPr.rFonts
        for attribut in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            schriften.attrib.pop(qn(attribut), None)
        stil.font.size = Pt(groesse)
        stil.font.bold = True
        stil.font.color.rgb = TEXT
        stil.paragraph_format.space_before = Pt(vorher)
        stil.paragraph_format.space_after = Pt(nachher)
        stil.paragraph_format.keep_with_next = True

    beschriftung = doc.styles["Caption"]
    beschriftung.font.name = SCHRIFT
    beschriftung.font.size = Pt(9)
    beschriftung.font.bold = False
    beschriftung.font.italic = False
    beschriftung.font.color.rgb = GEDAEMPFT
    beschriftung.paragraph_format.space_after = Pt(12)

    doc.styles["List Bullet"].font.name = SCHRIFT


# ---------------------------------------------------------------- Textauszeichnung

AUSZEICHNUNG = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def text_mit_auszeichnung(absatz, text: str) -> None:
    for teil in AUSZEICHNUNG.split(text):
        if not teil:
            continue
        if teil.startswith("**"):
            absatz.add_run(teil[2:-2]).bold = True
        elif teil.startswith("`"):
            lauf = absatz.add_run(teil[1:-1])
            lauf.font.name = "Consolas"
            lauf.font.size = Pt(9.5)
        else:
            absatz.add_run(teil)


def spaltenbreiten(zeilen: list[list[str]]) -> list[Cm]:
    """Jede Spalte bekommt mindestens Platz für ihr längstes Wort, der Rest
    wird nach Textmenge verteilt."""
    zeichen_cm = 0.19  # mittlere Zeichenbreite bei Arial 9 pt
    minima, gewichte = [], []
    for spalte in zip(*zeilen):
        laengstes_wort = max(len(wort) for inhalt in spalte for wort in (inhalt.split() or [""]))
        minima.append(laengstes_wort * zeichen_cm + 0.45)
        gewichte.append(sum(len(inhalt) for inhalt in spalte) / len(spalte) + 4)
    minima = [m * min(1, TEXTBREITE_CM / sum(minima)) for m in minima]
    rest = TEXTBREITE_CM - sum(minima)
    return [Cm(m + rest * g / sum(gewichte)) for m, g in zip(minima, gewichte)]


# ---------------------------------------------------------------- Bausteine


class Dokument:
    def __init__(self, meta: dict[str, str]):
        self.meta = meta
        self.doc = Document()
        vorlagen(self.doc)
        abschnitt = self.doc.sections[0]
        abschnitt.page_width, abschnitt.page_height = Cm(21), Cm(29.7)
        abschnitt.left_margin, abschnitt.right_margin = Cm(3), Cm(2.5)
        abschnitt.top_margin, abschnitt.bottom_margin = Cm(2.5), Cm(2.5)
        self.kapitel = 0
        self.unterkapitel = 0
        self.erstes_kapitel_im_abschnitt = True
        self.tabellenbeschriftung: str | None = None
        self.in_quellen = False

    # Deckblatt ---------------------------------------------------------------

    def deckblatt(self) -> None:
        m = self.meta
        doc = self.doc

        def zeile(text, groesse=11, fett=False, farbe=TEXT, abstand=6):
            absatz = doc.add_paragraph()
            absatz.alignment = WD_ALIGN_PARAGRAPH.CENTER
            absatz.paragraph_format.space_after = Pt(abstand)
            lauf = absatz.add_run(text)
            lauf.font.size = Pt(groesse)
            lauf.bold = fett
            lauf.font.color.rgb = farbe
            if text.startswith("["):
                lauf.font.highlight_color = WD_COLOR_INDEX.YELLOW
            return absatz

        zeile(m["hochschule"], 13, True, abstand=120)
        zeile("Projektdokumentation", 12, False, GEDAEMPFT, 18)
        zeile(m["titel"], 22, True, abstand=18)
        zeile(m["untertitel"], 12, False, GEDAEMPFT, 120)

        tabelle = doc.add_table(rows=0, cols=2)
        tabelle.alignment = WD_TABLE_ALIGNMENT.CENTER
        for beschriftung, schluessel in (
            ("Verfasser", "autor"),
            ("Studiengang", "studiengang"),
            ("Modul", "modul"),
            ("Betreuung", "betreuung"),
            ("Repository", "repository"),
        ):
            links, rechts = tabelle.add_row().cells
            links.width, rechts.width = Cm(4), Cm(9)
            links.paragraphs[0].add_run(beschriftung + ":").bold = True
            lauf = rechts.paragraphs[0].add_run(m[schluessel])
            if m[schluessel].startswith("["):
                lauf.font.highlight_color = WD_COLOR_INDEX.YELLOW

        zeile("", abstand=60)
        zeile(m["ort_datum"], 11)

    # Abschnitte und Seiten -------------------------------------------------

    def neuer_abschnitt(self, format_: str, kopftext: str) -> None:
        abschnitt = self.doc.add_section(WD_SECTION.NEW_PAGE)
        seitenzahlen(abschnitt, format_)
        for teil in (abschnitt.header, abschnitt.footer):
            teil.is_linked_to_previous = False
            for absatz in teil.paragraphs:
                absatz.text = ""

        kopf = abschnitt.header.paragraphs[0]
        kopf.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        lauf = kopf.add_run(kopftext)
        lauf.font.size = Pt(8.5)
        lauf.font.color.rgb = GEDAEMPFT

        fuss = abschnitt.footer.paragraphs[0]
        fuss.alignment = WD_ALIGN_PARAGRAPH.CENTER
        feld(fuss, "PAGE", "1")
        self.erstes_kapitel_im_abschnitt = True

    def seitenumbruch(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Inhalte ---------------------------------------------------------------

    def ueberschrift(self, ebene: int, text: str) -> None:
        nummeriert = not text.endswith("{-}")
        text = text.replace("{-}", "").strip()
        if ebene == 1:
            self.in_quellen = False
            if not self.erstes_kapitel_im_abschnitt:
                self.seitenumbruch()
            self.erstes_kapitel_im_abschnitt = False
            if nummeriert:
                self.kapitel += 1
                self.unterkapitel = 0
                text = f"{self.kapitel} {text}"
        elif nummeriert:
            self.unterkapitel += 1
            text = f"{self.kapitel}.{self.unterkapitel} {text}"
        self.doc.add_heading(text, level=ebene)

    def absatz(self, text: str) -> None:
        absatz = self.doc.add_paragraph()
        if self.in_quellen and text.startswith("["):
            absatz.paragraph_format.left_indent = Cm(1)
            absatz.paragraph_format.first_line_indent = Cm(-1)
            absatz.paragraph_format.line_spacing = 1.15
            absatz.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            absatz.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        text_mit_auszeichnung(absatz, text)

    def aufzaehlung(self, text: str, nummer: str | None) -> None:
        if nummer is None:
            absatz = self.doc.add_paragraph(style="List Bullet")
        else:
            absatz = self.doc.add_paragraph()
            absatz.paragraph_format.left_indent = Cm(0.75)
            absatz.paragraph_format.first_line_indent = Cm(-0.75)
            absatz.add_run(f"{nummer}\t")
        absatz.paragraph_format.space_after = Pt(3)
        text_mit_auszeichnung(absatz, text)

    def abbildung(self, beschriftung: str, pfad: Path) -> None:
        with Image.open(pfad) as bild:
            breite, hoehe = bild.size
        breite_cm = min(TEXTBREITE_CM, MAX_BILDHOEHE_CM * breite / hoehe)
        absatz = self.doc.add_paragraph()
        absatz.alignment = WD_ALIGN_PARAGRAPH.CENTER
        absatz.paragraph_format.space_before = Pt(6)
        nicht_trennen(absatz)
        absatz.add_run().add_picture(str(pfad), width=Cm(breite_cm))

        unter = self.doc.add_paragraph(style="Caption")
        unter.alignment = WD_ALIGN_PARAGRAPH.CENTER
        unter.add_run("Abbildung ")
        feld(unter, "SEQ Abbildung \\* ARABIC", "1")
        unter.paragraph_format.space_after = Pt(0)
        unter.add_run(f": {beschriftung}")
        # Quelle als eigener Absatz, damit sie nicht ins Abbildungsverzeichnis wandert
        quelle = self.doc.add_paragraph()
        quelle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        quelle.paragraph_format.space_after = Pt(12)
        lauf = quelle.add_run("Quelle: eigene Darstellung")
        lauf.font.size = Pt(8)
        lauf.font.color.rgb = GEDAEMPFT

    def tabelle(self, zeilen: list[list[str]]) -> None:
        if self.tabellenbeschriftung:
            ueber = self.doc.add_paragraph(style="Caption")
            ueber.paragraph_format.space_after = Pt(4)
            nicht_trennen(ueber)
            ueber.add_run("Tabelle ")
            feld(ueber, "SEQ Tabelle \\* ARABIC", "1")
            ueber.add_run(f": {self.tabellenbeschriftung}")
            self.tabellenbeschriftung = None

        tabelle = self.doc.add_table(rows=len(zeilen), cols=len(zeilen[0]))
        tabellenrahmen(tabelle)
        tabelle.autofit = False
        breiten = spaltenbreiten(zeilen)
        for i, zeile in enumerate(zeilen):
            for j, inhalt in enumerate(zeile):
                zelle = tabelle.cell(i, j)
                zelle.width = breiten[j]
                absatz = zelle.paragraphs[0]
                absatz.paragraph_format.space_after = Pt(0)
                absatz.paragraph_format.line_spacing = 1.1
                text_mit_auszeichnung(absatz, inhalt)
                for lauf in absatz.runs:
                    lauf.font.size = Pt(9)
                    if i == 0:
                        lauf.bold = True
                if i == 0:
                    zellfarbe(zelle, KOPF_HINTERGRUND)
        kopfzeile_wiederholen(tabelle.rows[0])
        for zeile in tabelle.rows:
            zeile_zusammenhalten(zeile)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def verzeichnis(self, art: str) -> None:
        titel, anweisung = {
            "toc": ("Inhaltsverzeichnis", 'TOC \\o "1-2" \\h \\z \\u'),
            "lof": ("Abbildungsverzeichnis", 'TOC \\h \\z \\c "Abbildung"'),
            "lot": ("Tabellenverzeichnis", 'TOC \\h \\z \\c "Tabelle"'),
        }[art]
        if not self.erstes_kapitel_im_abschnitt:
            self.seitenumbruch()
        self.erstes_kapitel_im_abschnitt = False
        # Eigene Formatierung statt Heading, damit das Verzeichnis sich nicht selbst aufführt
        kopf = self.doc.add_paragraph()
        kopf.paragraph_format.space_after = Pt(12)
        lauf = kopf.add_run(titel)
        lauf.bold = True
        lauf.font.size = Pt(16)
        feld(self.doc.add_paragraph(), anweisung, "Verzeichnis wird beim Öffnen in Word aktualisiert (F9).")


# ---------------------------------------------------------------- Einlesen


def metadaten(text: str) -> tuple[dict[str, str], str]:
    _, kopf, rumpf = text.split("---", 2)
    meta = {}
    for zeile in kopf.strip().splitlines():
        schluessel, wert = zeile.split(":", 1)
        meta[schluessel.strip()] = wert.strip().strip('"')
    return meta, rumpf


def tabellenzeile(zeile: str) -> list[str]:
    return [zelle.strip() for zelle in zeile.strip().strip("|").split("|")]


def bauen(ziel: Path = ZIEL) -> Path:
    meta, rumpf = metadaten(QUELLE.read_text(encoding="utf-8"))
    d = Dokument(meta)
    d.deckblatt()
    kopftext = meta["titel"]

    zeilen = rumpf.splitlines()
    i = 0
    while i < len(zeilen):
        zeile = zeilen[i].rstrip()
        i += 1
        if not zeile:
            continue

        if marker := re.fullmatch(r"<!-- (\w+) -->", zeile):
            art = marker.group(1)
            if art == "frontmatter":
                d.neuer_abschnitt("lowerRoman", kopftext)
            elif art == "main":
                d.neuer_abschnitt("decimal", kopftext)
            elif art == "quellen":
                d.in_quellen = True
            else:
                d.verzeichnis(art)
        elif zeile.startswith("## "):
            d.ueberschrift(2, zeile[3:])
        elif zeile.startswith("# "):
            d.ueberschrift(1, zeile[2:])
        elif bild := re.fullmatch(r"!\[(.+)\]\((.+)\)", zeile):
            d.abbildung(bild.group(1), (HIER / bild.group(2)).resolve())
        elif zeile.startswith("Tabelle: "):
            d.tabellenbeschriftung = zeile[len("Tabelle: "):]
        elif zeile.startswith("|"):
            block = [zeile]
            while i < len(zeilen) and zeilen[i].startswith("|"):
                block.append(zeilen[i])
                i += 1
            d.tabelle([tabellenzeile(z) for z in block if not re.fullmatch(r"\|[\s\-|]+\|", z.strip())])
        elif zeile.startswith("- "):
            d.aufzaehlung(zeile[2:], None)
        elif nummer := re.match(r"(\d+\.) (.+)", zeile):
            d.aufzaehlung(nummer.group(2), nummer.group(1))
        else:
            d.absatz(zeile)

    d.doc.save(ziel)
    return ziel


# ---------------------------------------------------------------- Word: Felder und PDF

WORD_SKRIPT = r"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {
  $doc = $word.Documents.Open('{docx}')
  # Erst die Nummern (SEQ = Typ 12), dann zweimal die Verzeichnisse,
  # damit auch deren Seitenzahlen stimmen
  foreach ($f in $doc.Fields) { if ($f.Type -eq 12) { $f.Update() | Out-Null } }
  1..2 | ForEach-Object {
    foreach ($v in $doc.TablesOfContents) { $v.Update() }
    foreach ($v in $doc.TablesOfFigures) { $v.Update() }
  }
  $doc.Save()
  $doc.SaveAs([ref]'{pdf}', [ref]17)
  $doc.Close()
} finally { $word.Quit() }
"""


def mit_word_aktualisieren(docx: Path) -> None:
    if sys.platform != "win32" or not shutil.which("powershell"):
        print("Word nicht verfügbar: Verzeichnisse beim Öffnen mit F9 aktualisieren.")
        return
    skript = WORD_SKRIPT.replace("{docx}", str(docx)).replace("{pdf}", str(docx.with_suffix(".pdf")))
    ergebnis = subprocess.run(["powershell", "-NoProfile", "-Command", skript], capture_output=True, text=True)
    if ergebnis.returncode != 0:
        print("Word-Aktualisierung fehlgeschlagen:", ergebnis.stderr.strip())
    else:
        print("Verzeichnisse aktualisiert, PDF erzeugt:", docx.with_suffix(".pdf"))


if __name__ == "__main__":
    pfad = bauen(Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ZIEL)
    print("Word-Datei erzeugt:", pfad)
    mit_word_aktualisieren(pfad)
